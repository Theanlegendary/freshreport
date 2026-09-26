"""
agent_export.py
Module for exporting Metfone Express Agents in the official Annex_03 format.
Generates 3 sheets: Khmer, English, and Total matching Annex_03_Number_of_service_points_in_Phnom_Penh_and_the_provinces.xlsx.
"""

import os
import re
import json
import logging
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

log = logging.getLogger("agent_export")
HERE = os.path.dirname(os.path.abspath(__file__))

# ── 1. Province definitions & IDs ──
PROV_INFO = {
    'BAN': {'id': 1, 'en': 'Banteay Meanchey', 'kh': 'បន្ទាយមានជ័យ'},
    'BAT': {'id': 2, 'en': 'Battambang', 'kh': 'បាត់ដំបង'},
    'CHA': {'id': 3, 'en': 'Kampong Cham', 'kh': 'កំពង់ចាម'},
    'CHH': {'id': 4, 'en': 'Kampong Chhnang', 'kh': 'កំពង់ឆ្នាំង'},
    'SPE': {'id': 5, 'en': 'Kampong Speu', 'kh': 'កំពង់ស្ពឺ'},
    'THO': {'id': 6, 'en': 'Kampong Thom', 'kh': 'កំពង់ធំ'},
    'KAM': {'id': 7, 'en': 'Kampot', 'kh': 'កំពត'},
    'KAN': {'id': 8, 'en': 'Kandal', 'kh': 'កណ្តាល'},
    'KOH': {'id': 9, 'en': 'Koh Kong', 'kh': 'កោះកុង'},
    'KRA': {'id': 10, 'en': 'Kratie', 'kh': 'ក្រចេះ'},
    'MON': {'id': 11, 'en': 'Mondulkiri', 'kh': 'មណ្ឌលគិរី'},
    'PNP': {'id': 12, 'en': 'Phnom Penh', 'kh': 'រាជធានីភ្នំពេញ'},
    'PRH': {'id': 13, 'en': 'Preah Vihear', 'kh': 'ព្រះវិហារ'},
    'PRE': {'id': 14, 'en': 'Prey Veng', 'kh': 'ព្រៃវែង'},
    'PUR': {'id': 15, 'en': 'Pursat', 'kh': 'ពោធិ៍សាត់'},
    'ROT': {'id': 16, 'en': 'Ratanakiri', 'kh': 'រតនគិរី'},
    'SIE': {'id': 17, 'en': 'Siem Reap', 'kh': 'សៀមរាប'},
    'SIH': {'id': 18, 'en': 'Preah Sihanouk', 'kh': 'ព្រះសីហនុ'},
    'STU': {'id': 19, 'en': 'Stung Treng', 'kh': 'ស្ទឹងត្រែង'},
    'SVA': {'id': 20, 'en': 'Svay Rieng', 'kh': 'ស្វាយរៀង'},
    'TAK': {'id': 21, 'en': 'Takeo', 'kh': 'តាកែវ'},
    'ODD': {'id': 22, 'en': 'Oddar Meanchey', 'kh': 'ឧត្តរមានជ័យ'},
    'KEP': {'id': 23, 'en': 'Kep', 'kh': 'កែប'},
    'KEB': {'id': 23, 'en': 'Kep', 'kh': 'កែប'},
    'PAI': {'id': 24, 'en': 'Pailin', 'kh': 'ប៉ៃលិន'},
    'TBK': {'id': 25, 'en': 'Tboung Khmum', 'kh': 'ត្បូងឃ្មុំ'},
}

DISTRICT_FALLBACK_EN = {
    "PRE": "Prey Veng", "PNP": "Chamkar Mon", "SVA": "Svay Rieng",
    "KAN": "Ta Khmau", "KAM": "Kampot", "KOH": "Khemarak Phoumin",
    "SIH": "Preah Sihanouk", "SPE": "Chbar Mon", "TAK": "Doun Kaev",
    "BAN": "Serei Saophoan", "BAT": "Battambang", "CHH": "Kampong Chhnang",
    "PUR": "Pursat", "SIE": "Siem Reap", "PRH": "Preah Vihear",
    "ODD": "Samraong", "THO": "Steung Saen", "CHA": "Kampong Cham",
    "KRA": "Kratie", "TBK": "Suong", "ROT": "Banlung",
    "MON": "Senmonorom", "STU": "Stung Treng", "KEP": "Kep", "PAI": "Pailin"
}

DISTRICT_FALLBACK_KH = {
    "PRE": "ព្រៃវែង", "PNP": "ចំការមន", "SVA": "ស្វាយរៀង",
    "KAN": "តាខ្មៅ", "KAM": "កំពត", "KOH": "ខេមរភូមិន្ទ",
    "SIH": "ព្រះសីហនុ", "SPE": "ច្បារមន", "TAK": "ដូនកែវ",
    "BAN": "សិរីសោភ័ណ", "BAT": "បាត់ដំបង", "CHH": "កំពង់ឆ្នាំង",
    "PUR": "ពោធិ៍សាត់", "SIE": "សៀមរាប", "PRH": "ព្រះវិហារ",
    "ODD": "សំរោង", "THO": "ស្ទឹងសែន", "CHA": "កំពង់ចាម",
    "KRA": "ក្រចេះ", "TBK": "សួង", "ROT": "បានលុង",
    "MON": "សែនមនោរម្យ", "STU": "ស្ទឹងត្រែង", "KEP": "កែប", "PAI": "ប៉ៃលិន"
}

_NIS_CACHE = None
_WC_CACHE = None

def _norm_str(s):
    return re.sub(r'[^a-zA-Z0-9\u1780-\u17FF]', '', str(s or '').lower())

def get_nis_gazetteer():
    global _NIS_CACHE
    if _NIS_CACHE is not None:
        return _NIS_CACHE
    nis_path = os.path.join(HERE, "cambodia_gazetteer_nis.json")
    if os.path.exists(nis_path):
        try:
            with open(nis_path, "r", encoding="utf-8") as f:
                _NIS_CACHE = json.load(f)
                return _NIS_CACHE
        except Exception:
            pass
    _NIS_CACHE = []
    return _NIS_CACHE

def get_ward_commune_map():
    global _WC_CACHE
    if _WC_CACHE is not None:
        return _WC_CACHE
    _WC_CACHE = {}
    path = os.path.join(HERE, "Địa chỉ đơn vị hành chính tại campuchia 2025.xlsx")
    if os.path.exists(path):
        try:
            df_wc = pd.read_excel(path, sheet_name='Ward and Commune')
            for _, r in df_wc.iterrows():
                po = str(r.get('POST OFFICE', '')).strip().upper()
                if po and po != 'NAN':
                    _WC_CACHE[po] = {
                        'prov_en': str(r.iloc[3] or '').strip(),
                        'prov_kh': str(r.iloc[4] or '').strip(),
                        'dist_en': str(r.iloc[7] or '').strip(),
                        'dist_kh': str(r.iloc[8] or '').strip(),
                        'comm_en': str(r.iloc[11] or '').strip(),
                        'comm_kh': str(r.iloc[12] or '').strip(),
                        'manager': str(r.get('MANAGER NAME', '')).strip(),
                        'phone': str(r.get('PHONE NUMBER', '')).strip(),
                    }
        except Exception as e:
            log.warning("Could not load Ward and Commune map: %s", e)
    return _WC_CACHE

def match_nis_division(branch_code, comm_en, comm_kh, dist_en=""):
    nis = get_nis_gazetteer()
    p_info = PROV_INFO.get(branch_code, {})
    p_id = p_info.get('id')
    
    c_en = _norm_str(comm_en)
    c_kh = _norm_str(comm_kh)
    d_en = _norm_str(dist_en)
    
    # 1. Exact match with district
    if d_en:
        for g in nis:
            if p_id and g['prov_id'] != p_id:
                continue
            if _norm_str(g['dist_en']) == d_en or _norm_str(g['dist_kh']) == d_en:
                if _norm_str(g['comm_en']) == c_en or (c_kh and _norm_str(g['comm_kh']) == c_kh):
                    return g

    # 2. Match commune within province
    for g in nis:
        if p_id and g['prov_id'] != p_id:
            continue
        if _norm_str(g['comm_en']) == c_en or (c_kh and _norm_str(g['comm_kh']) == c_kh):
            return g
            
    # 3. Substring match within province
    if len(c_en) >= 4:
        for g in nis:
            if p_id and g['prov_id'] != p_id:
                continue
            gc_en = _norm_str(g['comm_en'])
            if c_en in gc_en or gc_en in c_en:
                return g

    # 4. Match nationwide
    for g in nis:
        if _norm_str(g['comm_en']) == c_en or (c_kh and _norm_str(g['comm_kh']) == c_kh):
            return g

    return None

def strip_dept_code(value, code=""):
    text = str(value or "").strip()
    if not text or text.lower() == "nan":
        return ""
    if code:
        text = re.sub(rf"^{re.escape(str(code).strip())}\s*-\s*", "", text, flags=re.IGNORECASE)
    return re.sub(r"^[A-Z0-9]{3,10}\s*-\s*", "", text).strip()

def build_agents_dataframe(post_offices=None, branch_filter=None):
    """
    Builds a cleaned DataFrame of agents with complete administrative details.
    Uses post_offices list if provided; otherwise loads from pickup_branch_lookup.csv.
    """
    if post_offices is None or len(post_offices) == 0:
        pk_path = os.path.join(HERE, "pickup_branch_lookup.csv")
        if os.path.exists(pk_path):
            df_raw = pd.read_csv(pk_path)
            post_offices = df_raw.to_dict(orient="records")
        else:
            post_offices = []

    wc_map = get_ward_commune_map()

    # Normalize branch_filter
    filter_branches = set()
    if branch_filter:
        if isinstance(branch_filter, str):
            filter_branches = {b.strip().upper() for b in branch_filter.replace(";", ",").split(",") if b.strip()}
        elif isinstance(branch_filter, (list, tuple, set)):
            filter_branches = {str(b).strip().upper() for b in branch_filter if str(b).strip()}
        filter_branches.discard("ALL")
        filter_branches.discard("AGENT")
        filter_branches.discard("AGENTS")

    rows = []
    seen_codes = set()

    for item in post_offices:
        if not isinstance(item, dict):
            continue
            
        code = str(item.get("code") or item.get("Post Code") or item.get("Pickup Branch") or item.get("Post code") or "").strip().upper()
        if not code or code in seen_codes:
            continue

        # Facility classification: MUST BE AGENT
        typ = str(item.get("type") or item.get("Type") or item.get("Post Office Level") or item.get("Category") or "").strip().upper()
        is_agent = (
            typ == "AGENT"
            or "AGENT" in typ
            or "DEALER" in typ
            or bool(re.search(r"^[A-Z]{3}A\d+$", code))
        )
        if not is_agent:
            continue

        # Branch Code
        branch = item.get("branch") if isinstance(item.get("branch"), dict) else {}
        branch_code = str(
            item.get("parentDepartmentCode")
            or item.get("Branch Code")
            or item.get("Branch")
            or branch.get("code")
            or (code[:3] if len(code) >= 3 else "")
        ).strip().upper()
        # Handle "BAN - Banteay Meanchey" format in Branch
        if " - " in branch_code:
            branch_code = branch_code.split(" - ")[0].strip()

        if filter_branches and branch_code not in filter_branches:
            continue

        seen_codes.add(code)

        # Name
        raw_name = item.get("name") or item.get("Post Office Name") or item.get("enUsName") or item.get("kmKhmName") or ""
        agent_name = strip_dept_code(raw_name, code)
        if not agent_name:
            agent_name = str(item.get("Commune") or item.get("Commune EN") or code).strip()

        # Coordinates
        lat = item.get("latitude") or item.get("Latitude") or item.get("Lat (*)") or None
        lon = item.get("longitude") or item.get("Longitude") or item.get("Long (*)") or None
        try:
            lat = round(float(lat), 6) if lat is not None and str(lat).strip() not in ("", "nan", "None") else None
        except Exception:
            lat = None
        try:
            lon = round(float(lon), 6) if lon is not None and str(lon).strip() not in ("", "nan", "None") else None
        except Exception:
            lon = None

        # Province info
        p_info = PROV_INFO.get(branch_code, {'id': None, 'en': branch_code, 'kh': branch_code})
        prov_id = p_info['id']
        prov_code = branch_code
        prov_en = p_info['en']
        prov_kh = p_info['kh']

        # District & Commune extraction
        comm_en_raw = strip_dept_code(item.get("Commune") or item.get("Commune EN") or agent_name, code)
        comm_kh_raw = strip_dept_code(item.get("Commune Khmer") or item.get("kmKhmName") or comm_en_raw, code)
        dist_en_raw = strip_dept_code(item.get("District") or "", code)

        dist_id = None
        comm_id = None
        dist_en = dist_en_raw or DISTRICT_FALLBACK_EN.get(branch_code, "")
        dist_kh = DISTRICT_FALLBACK_KH.get(branch_code, "")
        comm_en = comm_en_raw
        comm_kh = comm_kh_raw

        # Try match from Ward and Commune if code is in it
        if code in wc_map:
            wc = wc_map[code]
            if wc.get('dist_en'): dist_en = wc['dist_en']
            if wc.get('dist_kh'): dist_kh = wc['dist_kh']
            if wc.get('comm_en'): comm_en = wc['comm_en']
            if wc.get('comm_kh'): comm_kh = wc['comm_kh']

        # Match NIS
        nis_match = match_nis_division(branch_code, comm_en, comm_kh, dist_en)
        if nis_match:
            dist_id = nis_match.get("dist_id")
            dist_en = nis_match.get("dist_en") or dist_en
            dist_kh = nis_match.get("dist_kh") or dist_kh
            comm_id = nis_match.get("comm_id")
            comm_en = nis_match.get("comm_en") or comm_en
            comm_kh = nis_match.get("comm_kh") or comm_kh

        # Address details
        addr_en = item.get("Address") or f"{comm_en}, {dist_en}, {prov_en}, Cambodia"
        addr_kh = item.get("Address details_kh")
        if not addr_kh or str(addr_kh).strip() == "" or str(addr_kh).lower() == "nan":
            # Format standard Khmer address details
            kh_parts = []
            if comm_kh: kh_parts.append(f"ឃុំ/សង្កាត់ {comm_kh}")
            if dist_kh: kh_parts.append(f"ស្រុក/ខណ្ឌ {dist_kh}")
            if prov_kh: kh_parts.append(f"ខេត្ត/ក្រុង {prov_kh}")
            addr_kh = ", ".join(kh_parts) if kh_parts else comm_kh

        rows.append({
            "Express Code": code,
            "Express Name": agent_name,
            "province_id": prov_id,
            "province_code": prov_code,
            "province_en": prov_en,
            "province_kh": prov_kh,
            "district_id": dist_id,
            "district_en": dist_en,
            "district_kh": dist_kh,
            "commune_id": comm_id,
            "commune_en": comm_en,
            "commune_kh": comm_kh,
            "Address details_en": addr_en,
            "Address details_kh": addr_kh,
            "Lat (*)": lat,
            "Long (*)": lon,
            "Remark (Branch/Service Point)": "Agent",
        })

    df = pd.DataFrame(rows)
    if not df.empty:
        # Sort by province_id, then Express Code
        df["_sort_pid"] = df["province_id"].fillna(99).astype(int)
        df = df.sort_values(by=["_sort_pid", "Express Code"]).reset_index(drop=True)
        df.drop(columns=["_sort_pid"], inplace=True)
        # Add No. (1-indexed)
        df.insert(0, "No", range(1, len(df) + 1))

    return df


def generate_agent_export_excel(df: pd.DataFrame, out_path: str, title: str = None) -> str:
    """
    Creates an Excel workbook styled identically to Annex_03.
    Contains three sheets: Khmer, English, and Total.
    """
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    RED_HEADER_FILL = "FFC00000"
    WHITE_FONT = "FFFFFFFF"
    DARK_TEXT = "000000"
    BORDER_CLR = "CCCCCC"

    thin_border = Border(
        left=Side(style="thin", color=BORDER_CLR),
        right=Side(style="thin", color=BORDER_CLR),
        top=Side(style="thin", color=BORDER_CLR),
        bottom=Side(style="thin", color=BORDER_CLR),
    )

    header_font = Font(name="Times New Roman", size=10, bold=True, color=WHITE_FONT)
    header_fill = PatternFill(start_color=RED_HEADER_FILL, end_color=RED_HEADER_FILL, fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Times New Roman", size=10)
    data_align_left = Alignment(horizontal="left", vertical="center")
    data_align_center = Alignment(horizontal="center", vertical="center")

    title_text = (
        "បញ្ជីឈ្មោះភ្នាក់ងារ នៅរាជធានីភ្នំពេញនិងតាមបណ្តាខេត្ត\n"
        "List of Agents in Phnom Penh and Provinces"
    )

    # ── 1. SHEET: Khmer ──
    ws_kh = wb.create_sheet(title="Khmer")
    ws_kh.views.sheetView[0].showGridLines = True

    kh_columns = [
        ("No", "No\nលេខរៀង", 8.7),
        ("Express Code", "Express Code\nលេខកូដសាខា", 13.0),
        ("Express Name", "Express Name", 20.0),
        ("province_id", "province_id", 11.5),
        ("province_code", "province_code", 10.0),
        ("province_kh", "province_kh", 16.0),
        ("district_id", "district_id", 11.5),
        ("district_kh", "district_kh\nស្រុក/ ខណ្ឌ ជាភាសាខ្មែរ", 22.0),
        ("commune_id", "commune_id\nលេខកូដសម្គាលឃុំ /សង្កាត់", 14.0),
        ("commune_kh", "commune_kh\nឃុំ /សង្កាត់ជាភាសាខ្មែរ", 16.0),
        ("Address details_kh", "Address details_kh\nព័ត៌មានលម្អិតអាសយដ្ឋាន ជាភាសាខ្មែរ", 35.0),
        ("Lat (*)", "Lat (*)", 12.0),
        ("Long (*)", "Long (*)", 13.0),
        ("Remark (Branch/Service Point)", "Remark (Branch/Service Point)", 16.5),
    ]

    # Title row (Row 1)
    ws_kh.merge_cells("A1:N1")
    t_cell = ws_kh.cell(1, 1, value=title_text)
    t_cell.font = Font(name="Times New Roman", size=13, bold=True, color="990000")
    t_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_kh.row_dimensions[1].height = 50.0

    # Header row (Row 2)
    ws_kh.row_dimensions[2].height = 54.75
    for c_idx, (col_key, col_header, width) in enumerate(kh_columns, 1):
        cell = ws_kh.cell(2, c_idx, value=col_header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        col_letter = get_column_letter(c_idx)
        ws_kh.column_dimensions[col_letter].width = width

    # Data rows (Row 3+)
    for r_idx, row in df.iterrows():
        excel_row = r_idx + 3
        ws_kh.row_dimensions[excel_row].height = 24.0
        for c_idx, (col_key, _, _) in enumerate(kh_columns, 1):
            val = row.get(col_key)
            val = "" if pd.isna(val) or val is None else val
            cell = ws_kh.cell(excel_row, c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_key in ("No", "Express Code", "province_id", "province_code", "district_id", "commune_id", "Remark (Branch/Service Point)"):
                cell.alignment = data_align_center
            elif col_key in ("Lat (*)", "Long (*)"):
                cell.alignment = data_align_center
                if isinstance(val, (int, float)):
                    cell.number_format = "0.000000"
            else:
                cell.alignment = data_align_left

    ws_kh.freeze_panes = "A3"
    ws_kh.auto_filter.ref = f"A2:N{len(df)+2}"

    # ── 2. SHEET: English ──
    ws_en = wb.create_sheet(title="English")
    ws_en.views.sheetView[0].showGridLines = True

    en_columns = [
        ("No", "No ", 6.0),
        ("Express Code", "Express Code", 13.0),
        ("Express Name", "Express Name", 20.0),
        ("province_id", "province_id", 11.5),
        ("province_code", "province_code", 14.0),
        ("province_en", "province_en", 18.0),
        ("district_id", "district_id", 15.5),
        ("district_en", "district_en", 20.0),
        ("commune_id", "commune_id", 13.0),
        ("commune_en", "commune_en", 16.0),
        ("Address details_en", "Address details_en", 35.0),
        ("Lat (*)", "Lat (*)", 12.0),
        ("Long (*)", "Long (*)", 13.0),
        ("Remark (Branch/Service Point)", "Remark (Branch/Service Point)", 16.0),
    ]

    ws_en.merge_cells("A1:N1")
    t_cell_en = ws_en.cell(1, 1, value=title_text)
    t_cell_en.font = Font(name="Times New Roman", size=13, bold=True, color="990000")
    t_cell_en.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_en.row_dimensions[1].height = 50.0

    ws_en.row_dimensions[2].height = 41.25
    for c_idx, (col_key, col_header, width) in enumerate(en_columns, 1):
        cell = ws_en.cell(2, c_idx, value=col_header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        col_letter = get_column_letter(c_idx)
        ws_en.column_dimensions[col_letter].width = width

    for r_idx, row in df.iterrows():
        excel_row = r_idx + 3
        ws_en.row_dimensions[excel_row].height = 24.0
        for c_idx, (col_key, _, _) in enumerate(en_columns, 1):
            val = row.get(col_key)
            val = "" if pd.isna(val) or val is None else val
            cell = ws_en.cell(excel_row, c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_key in ("No", "Express Code", "province_id", "province_code", "district_id", "commune_id", "Remark (Branch/Service Point)"):
                cell.alignment = data_align_center
            elif col_key in ("Lat (*)", "Long (*)"):
                cell.alignment = data_align_center
                if isinstance(val, (int, float)):
                    cell.number_format = "0.000000"
            else:
                cell.alignment = data_align_left

    ws_en.freeze_panes = "A3"
    ws_en.auto_filter.ref = f"A2:N{len(df)+2}"

    # ── 3. SHEET: Total ──
    ws_tot = wb.create_sheet(title="Total")
    ws_tot.views.sheetView[0].showGridLines = True

    tot_columns = [
        ("No", "No ", 6.0),
        ("Express Code", "Express Code", 13.0),
        ("Express Name", "Express Name", 20.0),
        ("province_id", "province_id", 11.5),
        ("province_code", "province_code", 14.0),
        ("province_en", "province_en", 18.0),
        ("province_kh", "province_kh", 16.0),
        ("district_id", "district_id", 15.5),
        ("district_en", "district_en", 20.0),
        ("district_kh", "district_kh", 20.0),
        ("commune_id", "commune_id", 13.0),
        ("commune_en", "commune_en", 16.0),
        ("commune_kh", "commune_kh", 16.0),
        ("Address details_en", "Address details_en", 32.0),
        ("Address details_kh", "Address details_kh", 35.0),
        ("Lat (*)", "Lat (*)", 12.0),
        ("Long (*)", "Long (*)", 13.0),
        ("Remark (Branch/Service Point)", "Remark (Branch/Service Point)", 18.0),
    ]

    ws_tot.merge_cells("A1:R1")
    t_cell_tot = ws_tot.cell(1, 1, value=title_text)
    t_cell_tot.font = Font(name="Times New Roman", size=13, bold=True, color="990000")
    t_cell_tot.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_tot.row_dimensions[1].height = 50.0

    ws_tot.row_dimensions[2].height = 36.0
    for c_idx, (col_key, col_header, width) in enumerate(tot_columns, 1):
        cell = ws_tot.cell(2, c_idx, value=col_header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        col_letter = get_column_letter(c_idx)
        ws_tot.column_dimensions[col_letter].width = width

    for r_idx, row in df.iterrows():
        excel_row = r_idx + 3
        ws_tot.row_dimensions[excel_row].height = 24.0
        for c_idx, (col_key, _, _) in enumerate(tot_columns, 1):
            val = row.get(col_key)
            val = "" if pd.isna(val) or val is None else val
            cell = ws_tot.cell(excel_row, c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_key in ("No", "Express Code", "province_id", "province_code", "district_id", "commune_id", "Remark (Branch/Service Point)"):
                cell.alignment = data_align_center
            elif col_key in ("Lat (*)", "Long (*)"):
                cell.alignment = data_align_center
                if isinstance(val, (int, float)):
                    cell.number_format = "0.000000"
            else:
                cell.alignment = data_align_left

    ws_tot.freeze_panes = "A3"
    ws_tot.auto_filter.ref = f"A2:R{len(df)+2}"

    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    print("Testing build_agents_dataframe...")
    df_agents = build_agents_dataframe()
    print("Built agents dataframe:", df_agents.shape)
    print(df_agents.head(3))
    out_file = os.path.join(HERE, "Annex_03_Agents_Phnom_Penh_and_Provinces.xlsx")
    generate_agent_export_excel(df_agents, out_file)
    print(f"Generated: {out_file} (Size: {os.path.getsize(out_file)} bytes)")
